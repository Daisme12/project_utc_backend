#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script to convert ProjectUTC README.md to a professional academic Microsoft Word (.docx) report.
Compliant with TCVN / University of Transport and Communications (UTC) thesis/report format.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), val)
            border.set(qn('w:sz'), sz)
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), color)
            tblBorders.append(border)
        tblPr[0].append(tblBorders)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def parse_inline(paragraph, text, base_font="Times New Roman", base_size=13, base_color=(0, 0, 0), is_italic_base=False):
    """
    Parses bold (**), italic (*), code (`), link ([text](url)) in markdown text
    and adds runs to paragraph.
    """
    tokens = re.split(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))', text)
    for tok in tokens:
        if not tok:
            continue
        
        # Bold
        if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            content = tok[2:-2]
            run = paragraph.add_run(content)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.bold = True
            run.font.italic = is_italic_base
            run.font.color.rgb = RGBColor(*base_color)
        
        # Italic
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
            content = tok[1:-1]
            run = paragraph.add_run(content)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.italic = True
            run.font.color.rgb = RGBColor(*base_color)
            
        # Code inline
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            content = tok[1:-1]
            run = paragraph.add_run(content)
            run.font.name = 'Consolas'
            run.font.size = Pt(base_size - 1.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(180, 40, 70)  # Dark pink / red code color
            
        # Link [label](url)
        elif tok.startswith('[') and ']' in tok and '(' in tok and tok.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', tok)
            if m:
                label, url = m.groups()
                run = paragraph.add_run(label)
                run.font.name = base_font
                run.font.size = Pt(base_size)
                run.font.underline = True
                run.font.color.rgb = RGBColor(30, 80, 160)
            else:
                run = paragraph.add_run(tok)
                run.font.name = base_font
                run.font.size = Pt(base_size)
        else:
            run = paragraph.add_run(tok)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.italic = is_italic_base
            run.font.color.rgb = RGBColor(*base_color)

def create_cover_page(doc):
    """
    Creates a formal academic cover page for UTC BTL/Đồ án.
    """
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_before = Pt(0)
    p_uni.paragraph_format.space_after = Pt(2)
    run1 = p_uni.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    run1.font.name = "Times New Roman"
    run1.font.size = Pt(13)
    run1.font.bold = False
    
    run2 = p_uni.add_run("TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI\n")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(14)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(16, 44, 87)
    
    run3 = p_uni.add_run("KHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN KỸ THUẬT PHẦN MỀM\n")
    run3.font.name = "Times New Roman"
    run3.font.size = Pt(13)
    run3.font.bold = True
    
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(40)
    run_line = p_line.add_run("━━━━━━━━━━━━━━━━━━━ ◆ ━━━━━━━━━━━━━━━━━━━")
    run_line.font.name = "Times New Roman"
    run_line.font.size = Pt(12)
    run_line.font.color.rgb = RGBColor(100, 100, 100)
    
    p_doc = doc.add_paragraph()
    p_doc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_doc.paragraph_format.space_after = Pt(20)
    run_doc = p_doc.add_run("BÁO CÁO BÀI TẬP LỚN / ĐỒ ÁN MÔN HỌC")
    run_doc.font.name = "Times New Roman"
    run_doc.font.size = Pt(18)
    run_doc.font.bold = True
    run_doc.font.color.rgb = RGBColor(190, 30, 45) # Deep red / Academic crimson
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    run_sub = p_sub.add_run("HỌC PHẦN: PHÁT TRIỂN ỨNG DỤNG DOANH NGHIỆP & HỆ THỐNG PHẦN MỀM")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(60)
    run_title = p_title.add_run("ĐỀ TÀI:\nXÂY DỰNG HỆ THỐNG QUẢN TRỊ & BÁN HÀNG\nTRỰC TUYẾN UTC (PROJECT UTC)")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)
    
    # Table info
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Cm(5.5), Cm(9.5)]
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = col_widths[idx]
            
    info_data = [
        ("Sinh viên thực hiện:", "Đoàn Mạnh"),
        ("Mã số sinh viên (MSSV):", "233030594"),
        ("Khoa / Chuyên ngành:", "Công nghệ Thông tin - Kỹ thuật Phần mềm"),
        ("Cơ sở đào tạo:", "Trường Đại học Giao thông Vận tải (UTC)"),
    ]
    
    for row_idx, (lbl, val) in enumerate(info_data):
        cell_lbl = table.cell(row_idx, 0)
        cell_val = table.cell(row_idx, 1)
        
        p_l = cell_lbl.paragraphs[0]
        p_l.paragraph_format.space_after = Pt(4)
        run_l = p_l.add_run(lbl)
        run_l.font.name = "Times New Roman"
        run_l.font.size = Pt(13)
        run_l.font.bold = True
        
        p_v = cell_val.paragraphs[0]
        p_v.paragraph_format.space_after = Pt(4)
        run_v = p_v.add_run(val)
        run_v.font.name = "Times New Roman"
        run_v.font.size = Pt(13)
        run_v.font.bold = (row_idx < 2) # bold student name & MSSV
        if row_idx == 0:
            run_v.font.color.rgb = RGBColor(16, 44, 87)
            
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(80)
    
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.paragraph_format.space_after = Pt(0)
    run_f = p_footer.add_run("HÀ NỘI, NĂM 2026")
    run_f.font.name = "Times New Roman"
    run_f.font.size = Pt(13)
    run_f.font.bold = True
    
    doc.add_page_break()

def convert_markdown_to_docx(md_path, docx_path):
    print(f"Reading markdown from: {md_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = docx.Document()

    # Set Standard Page Margins (UTC Thesis Standard: Left 3cm, Right 2cm, Top 2cm, Bottom 2cm)
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        
        # Configure Header & Footer
        header = section.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.paragraph_format.space_after = Pt(0)
        r_head = p_head.add_run("Báo cáo BTL: Hệ thống Quản trị & Bán hàng Trực tuyến UTC (Project UTC)")
        r_head.font.name = "Times New Roman"
        r_head.font.size = Pt(9)
        r_head.font.italic = True
        r_head.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot.paragraph_format.space_after = Pt(0)
        r_foot_label = p_foot.add_run("Trang ")
        r_foot_label.font.name = "Times New Roman"
        r_foot_label.font.size = Pt(10)
        r_foot_label.font.color.rgb = RGBColor(100, 100, 100)
        add_page_number(p_foot.add_run())

    # Build Cover Page
    create_cover_page(doc)

    lines = md_text.split('\n')
    i = 0
    total_lines = len(lines)

    # State variables
    in_code_block = False
    code_lang = ""
    code_lines = []

    in_table = False
    table_rows = []

    while i < total_lines:
        line = lines[i]
        stripped = line.strip()

        # Handle Fenced Code Blocks (```)
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_lines = []
                i += 1
                continue
            else:
                # End of code block -> Output as formatted Word block
                in_code_block = False
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                cell.width = Cm(16.0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # Header label for code block
                p_cell = cell.paragraphs[0]
                p_cell.paragraph_format.line_spacing = 1.15
                p_cell.paragraph_format.space_after = Pt(2)
                p_cell.paragraph_format.space_before = Pt(0)
                
                lang_label = f"Mã nguồn / Sơ đồ ({code_lang.upper()}):" if code_lang else "Mã nguồn:"
                r_lang = p_cell.add_run(lang_label + "\n")
                r_lang.font.name = "Consolas"
                r_lang.font.size = Pt(9.5)
                r_lang.font.bold = True
                r_lang.font.color.rgb = RGBColor(80, 100, 120)
                
                code_content = "\n".join(code_lines)
                r_code = p_cell.add_run(code_content)
                r_code.font.name = "Consolas"
                r_code.font.size = Pt(9.5)
                r_code.font.color.rgb = RGBColor(30, 41, 59)
                
                set_table_borders(tbl, color="CBD5E1", sz="4", val="single")
                
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_before = Pt(0)
                p_after.paragraph_format.space_after = Pt(6)
                
                code_lines = []
                code_lang = ""
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Markdown Tables (| ... |)
        if stripped.startswith('|') and stripped.endswith('|'):
            # Collect all contiguous table rows
            table_lines = []
            while i < total_lines and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table lines
            parsed_rows = []
            for tline in table_lines:
                # check if separator row like | :--- | :---: |
                cells = [c.strip() for c in tline.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', c) for c in cells):
                    continue # separator row
                parsed_rows.append(cells)
                
            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="CBD5E1", sz="4", val="single")
                
                for r_idx, row_cells in enumerate(parsed_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell = tbl.cell(r_idx, c_idx)
                        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        
                        text_val = row_cells[c_idx] if c_idx < len(row_cells) else ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.line_spacing = 1.2
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        
                        if is_header:
                            set_cell_background(cell, "1F497D") # Navy blue header
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            parse_inline(p, text_val, base_font="Times New Roman", base_size=11.5, base_color=(255, 255, 255))
                            for run in p.runs:
                                run.font.bold = True
                        else:
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "FFFFFF")
                            else:
                                set_cell_background(cell, "F8FAFC") # zebra striping
                                
                            # Center align STT or short codes
                            if c_idx == 0 and len(text_val) <= 6:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            parse_inline(p, text_val, base_font="Times New Roman", base_size=11, base_color=(0, 0, 0))
                            
                p_tbl_after = doc.add_paragraph()
                p_tbl_after.paragraph_format.space_before = Pt(2)
                p_tbl_after.paragraph_format.space_after = Pt(8)
            continue

        # Skip horizontal rules
        if stripped in ('---', '***', '___'):
            i += 1
            continue

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            parse_inline(p, text, base_font="Times New Roman", base_size=16, base_color=(16, 44, 87))
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        if stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            parse_inline(p, text, base_font="Times New Roman", base_size=14, base_color=(31, 73, 125))
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        if stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            parse_inline(p, text, base_font="Times New Roman", base_size=13, base_color=(20, 20, 20))
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        if stripped.startswith('#### '):
            text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            parse_inline(p, text, base_font="Times New Roman", base_size=12.5, base_color=(40, 40, 40))
            for run in p.runs:
                run.font.bold = True
                run.font.italic = True
            i += 1
            continue

        if stripped.startswith('##### '):
            text = stripped[6:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            parse_inline(p, text, base_font="Times New Roman", base_size=12, base_color=(60, 60, 60))
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        # Blockquote (> ...)
        if stripped.startswith('> '):
            text = stripped[2:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Cm(16.0)
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=80, bottom=80, left=140, right=120)
            set_table_borders(tbl, color="3B82F6", sz="12", val="single") # Blue left/all border
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            parse_inline(p, text, base_font="Times New Roman", base_size=12.5, base_color=(30, 40, 50), is_italic_base=True)
            
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Bullet List (* or -)
        if stripped.startswith('* ') or stripped.startswith('- '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(3)
            
            bullet_run = p.add_run("•  ")
            bullet_run.font.name = "Times New Roman"
            bullet_run.font.size = Pt(13)
            bullet_run.font.bold = True
            bullet_run.font.color.rgb = RGBColor(31, 73, 125)
            
            parse_inline(p, text, base_font="Times New Roman", base_size=13)
            i += 1
            continue

        # Numbered List (1. , 2. )
        m_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m_num:
            num_str, text = m_num.groups()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(3)
            
            num_run = p.add_run(f"{num_str}. ")
            num_run.font.name = "Times New Roman"
            num_run.font.size = Pt(13)
            num_run.font.bold = True
            num_run.font.color.rgb = RGBColor(16, 44, 87)
            
            parse_inline(p, text, base_font="Times New Roman", base_size=13)
            i += 1
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parse_inline(p, line, base_font="Times New Roman", base_size=13)
        i += 1

    print(f"Saving formatted Word document to: {docx_path}")
    doc.save(docx_path)
    print("Done! Document successfully generated.")

if __name__ == '__main__':
    md_file = "/Users/doanmanh/Documents/Học_tập/ProjectUTC/README.md"
    docx_file = "/Users/doanmanh/Documents/Học_tập/ProjectUTC/BAO_CAO_PROJECT_UTC.docx"
    convert_markdown_to_docx(md_file, docx_file)
